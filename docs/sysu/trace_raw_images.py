"""Trace processed SYSU image SHA to raw files; write private source paths locally."""

import csv
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data/sysu"
PROCESSED = ROOT / "data/SYSU_processed"
OUTPUT = ROOT / "tmp/sysu-private"
SOURCES = {
    "level1": ("mild+normal.xlsx", "mild"),
    "level2": ("moderate.xlsx", "moderate"),
    "level3": ("severe.xlsx", "severe"),
}


def sha256(path):
    with path.open("rb") as file:
        return hashlib.file_digest(file, "sha256").hexdigest()


def normalized_name(value):
    name = str(value or "").strip().replace(" ", "").replace("　", "")
    for prefix in ("new_", "v", "删除_"):
        if name.startswith(prefix):
            name = name[len(prefix):]
    return name


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    rows = list(csv.DictReader((PROCESSED / "sysu_samples.csv").open(encoding="utf-8-sig")))
    names = {}
    for level, (filename, sheet_name) in SOURCES.items():
        workbook = load_workbook(RAW / level / filename, read_only=True, data_only=True)
        sheet = workbook[sheet_name]
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        name_column = [str(value or "").strip() for value in headers].index("name")
        for index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            names[f"{level}_{index:03d}"] = normalized_name(values[name_column])
        workbook.close()

    raw_by_sha = defaultdict(list)
    role_counts = Counter()
    for path in sorted(RAW.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}:
            continue
        relative = path.relative_to(RAW)
        role = ("selected_training" if path.name.lower() == "病损.jpg" else
                "unnumbered_other_format" if path.stem == "病损" else
                "numbered_or_other_candidate" if path.stem.startswith("病损") else "other_image")
        role_counts[role] += 1
        raw_by_sha[sha256(path)].append({"raw_path": str(path), "raw_level": relative.parts[0],
                                        "patient_directory": relative.parts[1] if len(relative.parts) > 2 else "",
                                        "role": role})

    trace = []
    processed_by_sha = defaultdict(list)
    for row in rows:
        if not row["image_path"]:
            continue
        digest = sha256(PROCESSED / row["image_path"])
        processed_by_sha[digest].append(row)
        level = row["sample_id"].split("_")[0]
        for match in raw_by_sha.get(digest, [{}]):
            trace.append({"sample_id": row["sample_id"], "label": row["label"], "image_sha256": digest,
                          **match, "same_level_as_table": match.get("raw_level") == level,
                          "directory_name_equals_table_name": bool(match.get("patient_directory")) and
                          normalized_name(match["patient_directory"]) == names[row["sample_id"]],
                          "directory_name_has_table_name_prefix": bool(match.get("patient_directory")) and
                          normalized_name(match["patient_directory"]).startswith(names[row["sample_id"]])})
    fields = ["sample_id", "label", "image_sha256", "raw_path", "raw_level", "patient_directory", "role",
              "same_level_as_table", "directory_name_equals_table_name", "directory_name_has_table_name_prefix"]
    with (OUTPUT / "processed-to-raw-image-sha-2026-10-08.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        writer.writerows(trace)

    duplicate_results = []
    clinical_fields = [key for key in rows[0] if key not in {"sample_id", "label", "image_path", "sex", "age"}]
    for digest, group in processed_by_sha.items():
        if len(group) < 2:
            continue
        ids = {row["sample_id"] for row in group}
        source_rows = [entry for entry in trace if entry["sample_id"] in ids]
        selected = [entry for entry in raw_by_sha.get(digest, []) if entry["role"] == "selected_training"]
        duplicate_results.append({
            "sample_ids": sorted(ids), "image_sha256": digest,
            "table_names_equal_after_normalization": len({names[sample_id] for sample_id in ids}) == 1,
            "selected_raw_matches": len(selected),
            "selected_raw_levels": sorted({entry["raw_level"] for entry in selected}),
            "selected_directory_names_equal": len({normalized_name(entry["patient_directory"]) for entry in selected}) == 1,
            "each_row_has_same_level_exact_name_selected_source": all(any(entry["sample_id"] == sample_id and
                    entry.get("role") == "selected_training" and entry["same_level_as_table"] and
                    entry["directory_name_equals_table_name"] for entry in source_rows) for sample_id in ids),
            "each_row_has_same_level_name_prefix_selected_source": all(any(entry["sample_id"] == sample_id and
                    entry.get("role") == "selected_training" and entry["same_level_as_table"] and
                    entry["directory_name_has_table_name_prefix"] for entry in source_rows) for sample_id in ids),
            "non_demographic_clinical_fields_equal": all(len({row[field] for row in group}) == 1 for field in clinical_fields),
        })
    unmatched_selected = []
    for digest, matches in raw_by_sha.items():
        if digest in processed_by_sha:
            continue
        for match in matches:
            if match["role"] != "selected_training":
                continue
            directory_name = normalized_name(match["patient_directory"])
            candidates = [row["sample_id"] for row in rows if row["sample_id"].startswith(match["raw_level"] + "_")
                          and (directory_name.startswith(names[row["sample_id"]]) or
                               names[row["sample_id"]].startswith(directory_name))]
            unmatched_selected.append({"raw_path": match["raw_path"], "image_sha256": digest,
                                       "raw_level": match["raw_level"], "candidate_sample_ids": "|".join(candidates)})
    with (OUTPUT / "selected-images-not-in-processed-2026-10-08.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=["raw_path", "image_sha256", "raw_level", "candidate_sample_ids"])
        writer.writeheader()
        writer.writerows(unmatched_selected)
    name_buckets = defaultdict(list)
    for row in rows:
        name_buckets[names[row["sample_id"]]].append(row)
    if "" in name_buckets:
        raise ValueError("Empty source name needs review before proposing an identity group")
    candidate_ids = {name: f"candidate_P{index:03d}" for index, name in enumerate(sorted(name_buckets), start=1)}
    with (OUTPUT / "name-candidate-groups-2026-10-08.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=["sample_id", "candidate_patient_id", "label", "same_name_record_count", "review_status"])
        writer.writeheader()
        for row in rows:
            name = names[row["sample_id"]]
            writer.writerow({"sample_id": row["sample_id"], "candidate_patient_id": candidate_ids[name],
                             "label": row["label"], "same_name_record_count": len(name_buckets[name]),
                             "review_status": "unverified_source_name_group"})
    with (OUTPUT / "same-name-review-2026-10-08.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=["candidate_patient_id", "sample_id", "label", "source_workbook", "source_sheet", "source_excel_row"])
        writer.writeheader()
        for name, group in name_buckets.items():
            if len(group) < 2:
                continue
            for row in group:
                level, row_number = row["sample_id"].split("_")
                filename, sheet = SOURCES[level]
                writer.writerow({"candidate_patient_id": candidate_ids[name], "sample_id": row["sample_id"],
                                 "label": row["label"], "source_workbook": str(RAW / level / filename),
                                 "source_sheet": sheet, "source_excel_row": int(row_number)})
    summary = {"raw_image_roles": dict(role_counts), "processed_images": sum(bool(row["image_path"]) for row in rows),
               "processed_images_with_selected_raw_match": sum(any(entry["sample_id"] == row["sample_id"] and
                       entry.get("role") == "selected_training" for entry in trace) for row in rows if row["image_path"]),
               "duplicate_groups": duplicate_results,
               "source_name_candidate_groups": len(name_buckets),
               "repeated_name_groups": sum(len(group) > 1 for group in name_buckets.values()),
               "repeated_name_records": sum(len(group) for group in name_buckets.values() if len(group) > 1),
               "selected_raw_files_not_in_processed": len(unmatched_selected),
               "unmatched_selected_candidate_counts": [len(entry["candidate_sample_ids"].split("|"))
                        if entry["candidate_sample_ids"] else 0 for entry in unmatched_selected],
               "identity_status": "Name/path agreement is source evidence; not a confirmed patient identity or label correction",
               "raw_data_modified": False, "numbered_candidates_used_for_training": False}
    (OUTPUT / "raw-image-sha-summary-2026-10-08.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
