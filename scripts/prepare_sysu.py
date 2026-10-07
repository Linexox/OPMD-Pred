"""Create the SYSU CSV and exact-stem lesion image directory."""

import argparse
import csv
import shutil
from pathlib import Path

from openpyxl import load_workbook


FIELDS = [
    "是否吸烟",
    "吸烟年限（年）",
    "吸烟量（支/天）",
    "是否饮酒",
    "饮酒年限(年)",
    "饮酒量(ml/天)",
    "是否嚼槟榔",
    "嚼槟榔年限(年)",
    "嚼槟榔量（颗/天）",
    "TCT意见",
    "DNA倍体分析结果",
    "甲基化总数",
]
OUTPUT_FIELDS = ["sample_id", "label", "image_path", *FIELDS]
MISSING = {"", "NAN", "nan", "NA", "N/A", "null", "NULL", "/", "#VALUE!"}


def clean(value):
    value = "" if value is None else str(value).strip()
    if value in MISSING:
        return ""
    if value == "几十年":
        return "30"
    return value


def read_rows(path, sheet_name):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = list(workbook[sheet_name].iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    if not {"name", *FIELDS} <= set(headers):
        raise ValueError(f"Required source columns missing in {path.name}/{sheet_name}")
    return [dict(zip(headers, row)) for row in values[1:] if any(row)]


def normalized_name(value):
    value = value.replace(" ", "").replace("　", "")
    for prefix in ("new_", "v", "删除_"):
        if value.startswith(prefix):
            value = value[len(prefix) :]
    return value


def image_for_row(level_dir, patient_name, used_dirs):
    name = normalized_name(patient_name)
    candidates = [
        directory
        for directory in level_dir.iterdir()
        if directory.is_dir()
        and directory not in used_dirs
        and (
            normalized_name(directory.name).startswith(name)
            or name.startswith(normalized_name(directory.name))
        )
    ]
    if not candidates:
        return None
    if len(candidates) != 1:
        raise ValueError("Ambiguous image-directory binding: review source identity before preparing data")
    directory = candidates[0]
    used_dirs.add(directory)
    images = [
        path
        for path in directory.iterdir()
        if path.is_file() and path.stem == "病损"
    ]
    if len(images) > 1:
        raise ValueError("Multiple exact-stem lesion images: review image selection before preparing data")
    return images[0] if images else None


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-root", type=Path, default=Path("data/sysu"))
    parser.add_argument("--output", type=Path, required=True, help="New directory; existing data is never overwritten")
    args = parser.parse_args()
    output = args.output
    if output.exists():
        raise FileExistsError("Use a new output directory and review it before switching training data")
    image_output = output / "images"
    output.mkdir(parents=True)
    image_output.mkdir(exist_ok=True)

    sources = [
        ("level1", "level1_mild_or_normal", "mild+normal.xlsx", "mild"),
        ("level2", "level2_moderate", "moderate.xlsx", "moderate"),
        ("level3", "level3_severe", "severe.xlsx", "severe"),
    ]
    rows = []
    for level, label, filename, sheet_name in sources:
        source = args.source_root / level
        used_dirs = set()
        for excel_row, source_row in enumerate(read_rows(source / filename, sheet_name), start=2):
            record = {
                "sample_id": f"{level}_{excel_row:03d}",
                "label": label,
                "image_path": "",
            }
            for field in FIELDS:
                record[field] = clean(source_row.get(field))
            image = image_for_row(source, source_row.get("name", ""), used_dirs)
            if image is not None:
                destination_name = f"{record['sample_id']}__病损.jpg"
                shutil.copy2(image, image_output / destination_name)
                record["image_path"] = f"images/{destination_name}"
            rows.append(record)

    with (output / "sysu_samples.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=OUTPUT_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"wrote {len(rows)} records and {len(list(image_output.iterdir()))} images")


if __name__ == "__main__":
    main()
