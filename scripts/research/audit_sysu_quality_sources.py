"""Read original sources for moderate-grade records and prepare blinded development-photo review."""
from __future__ import annotations

from collections import Counter
import csv
import hashlib
import json
from pathlib import Path
import random
import re
import zipfile

from openpyxl import load_workbook
import pdfplumber
import pypdfium2 as pdfium
from PIL import Image, ImageDraw, ImageFont, ImageOps

ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "data/processed/sysu/rebuilt-20261009"
OUT = ROOT / "outputs/experiments/sysu/moderate-quality-20261016/source-review"
RAW = (ROOT / "data/sysu").resolve()


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def snippets(text, pattern, width=70):
    return [text[max(0, m.start()-width):min(len(text), m.end()+width)] for m in re.finditer(pattern, text)]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    cv = read(INPUT / "development_nested_cv_v3.json")
    entries = {e["sample_id"]: e for e in cv["development_entries"]}
    ids = set(entries)
    assert len(ids) == 117 and not ids & set(cv["locked_previous_test_ids"] + cv["additional_related_locked_ids"])
    assert sha(INPUT / "sysu_samples.csv") == cv["csv_sha256"]
    with (INPUT / "sysu_samples.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = {row["sample_id"]: row for row in csv.DictReader(stream) if row["sample_id"] in ids}
    provenance = {r["sample_id"]: r for r in read(INPUT / "record_provenance.private.json") if r["sample_id"] in ids}
    inv = read(INPUT / "inventory.private.json")
    old = {r["sample_id"]: r for r in inv["records"] if r["sample_id"] in ids}
    moderate = sorted(sid for sid, entry in entries.items() if entry["label"] == 1)
    assert len(moderate) == 31
    (OUT / "quality-rubric.md").write_text(
        "# Review criteria frozen before viewing photos\n\n"
        "Scope: 75 frozen development photographs only; no locked test photos.\n"
        "Review with Q001-Q075 codes, randomized order, no grade or prediction shown.\n"
        "Inspect focus, broad over/underexposure, occlusion, framing, non-photo overlays, and visible oral tissue.\n"
        "Ordinary saliva highlights, clinically atypical appearance, negative tests, and missing modalities are not exclusion reasons.\n"
        "Review 518-pixel aspect-padded input; investigate flagged photos at native resolution.\n"
        "Categories: no_obvious_major_technical_issue / technical_issue_candidate / non_oral_image.\n"
        "Describe visible defects and uncertainty; no lesion diagnosis, target-site confirmation, or clinical readability certification.\n"
        "No records or images are excluded or modified during review. Final clinical usability remains unconfirmed.\n",
        encoding="utf-8")
    file_inventory, documents, errors = [], {}, []
    case_sources = []
    for sid in moderate:
        prov, source = provenance[sid], old[sid]
        folders = [Path(p).resolve() for p in prov["bound_source_folders"]]
        assert all(folder.is_relative_to(RAW) for folder in folders)
        files = sorted({path for folder in folders for path in folder.rglob("*") if path.is_file()})
        for path in files:
            key = str(path)
            fingerprint = sha(path)
            file_inventory.append(dict(sample_id=sid, path=key, sha256=fingerprint, suffix=path.suffix.lower()))
            if path.suffix.lower() == ".pdf" and key not in documents:
                try:
                    with pdfplumber.open(path) as pdf:
                        pages = [dict(page=i+1, text=page.extract_text() or "") for i, page in enumerate(pdf.pages)]
                    text = "\n".join(page["text"] for page in pages)
                    documents[key] = dict(sha256=fingerprint, pages=pages, text_characters=len(text),
                        histology_filename=bool(re.search("病理|组织|活检|病检", path.name)),
                        histology_text_candidates=snippets(text, "病理诊断|组织学诊断|镜下|组织病理|中度.{0,8}增生|轻中度|轻.中度"),
                        site_candidates=snippets(text, "取材部位|采样部位|病损部位|病变部位|标本来源|标本名称"))
                except Exception as exc:
                    errors.append(dict(path=key, error=repr(exc)))
            elif path.suffix.lower() in (".txt", ".csv") and key not in documents:
                text = path.read_text(encoding="utf-8-sig", errors="replace")
                documents[key] = dict(sha256=fingerprint, text=text,
                    histology_text_candidates=snippets(text, "病理诊断|组织学诊断|镜下|组织病理|中度.{0,8}增生|轻中度|轻.中度"))
        case_sources.append(dict(sample_id=sid, expected_grade=1, source_row=prov["source_table"],
            source_folders=[str(p) for p in folders], folder_grade_evidence=[p.name for p in folders],
            binding_basis=prov["binding_basis"], source_name_key=source["name_key"],
            files=[str(path) for path in files], selected_image_sha=prov["image_sha256"],
            dna_status=prov["dna_status"], frozen_dna=prov.get("dna")))
    # Primary source rows are fixed; supplementary-sheet matches are candidates only.
    workbook = inv["workbooks"]["level2"]
    path = Path(workbook["path"]).resolve()
    assert path.is_relative_to(RAW) and sha(path) == workbook["sha256"]
    wb = load_workbook(path, read_only=False, data_only=True)
    supplementary = []
    names = {old[sid]["name_key"]: sid for sid in moderate}
    for sheet in wb.worksheets:
        matches = []
        if sheet.title == workbook["sheet"]:
            row_numbers = {old[sid]["raw_excel_row"]: sid for sid in moderate}
            for rownum, sid in row_numbers.items():
                cells = [dict(coordinate=cell.coordinate, value=str(cell.value) if cell.value is not None else None,
                    comment=cell.comment.text if cell.comment else None) for cell in sheet[rownum]]
                matches.append(dict(sample_id=sid, row=rownum, cells=cells))
        else:
            for row in sheet.iter_rows():
                values = [str(cell.value).strip() for cell in row if cell.value is not None]
                candidate_ids = {names[value] for value in values if value in names}
                if candidate_ids:
                    matches.append(dict(candidate_sample_ids=sorted(candidate_ids), row=row[0].row,
                        cells=[dict(coordinate=cell.coordinate, value=str(cell.value) if cell.value is not None else None,
                            comment=cell.comment.text if cell.comment else None) for cell in row]))
        supplementary.append(dict(sheet=sheet.title, rows=sheet.max_row, columns=sheet.max_column,
                                  embedded_images=len(sheet._images), matches=matches))
    with zipfile.ZipFile(path) as archive:
        embedded = [name for name in archive.namelist() if name.startswith("xl/media/")]
    image_ids = sorted(sid for sid, e in entries.items() if e["image_sha256"])
    assert len(image_ids) == 75
    random.Random(2026101601).shuffle(image_ids)
    image_manifest = []
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 27)
    for index, sid in enumerate(image_ids, start=1):
        code = f"Q{index:03d}"
        image = (INPUT / rows[sid]["image_path"]).resolve()
        assert image.is_relative_to(INPUT.resolve()) and sha(image) == entries[sid]["image_sha256"]
        folders = [Path(p).resolve() for p in provenance[sid]["bound_source_folders"]]
        plain = [p for folder in folders for p in folder.rglob("*") if p.is_file()
                 and p.suffix.lower() in (".jpg", ".jpeg", ".png") and p.stem.strip() == "病损"]
        matches = [str(p) for p in plain if sha(p) == entries[sid]["image_sha256"]]
        with Image.open(image) as source:
            source.load()
            tile = Image.new("RGB", (600, 580), "white")
            resized = ImageOps.pad(source.convert("RGB"), (518, 518), method=Image.Resampling.BILINEAR,
                                   color=(124, 116, 104), centering=(.5, .5))
            tile.paste(resized, (41, 46))
            ImageDraw.Draw(tile).text((41, 6), code, font=font, fill="black")
            item = dict(code=code, sample_id=sid, label=entries[sid]["label"], cached_image_path=str(image),
                sha256=entries[sid]["image_sha256"], width=source.width, height=source.height,
                exif_orientation=source.getexif().get(274), selected_unnumbered_source_matches=matches,
                plain_candidates=[str(p) for p in plain])
        image_manifest.append(item)
        if (index-1) % 6 == 0:
            sheet = Image.new("RGB", (1800, 1160), "#dedede")
        position = (index-1) % 6
        sheet.paste(tile, ((position % 3) * 600, (position // 3) * 580))
        if position == 5 or index == len(image_ids):
            output = OUT / "photo-sheets" / f"sheet-{(index-1)//6+1:02d}.jpg"
            output.parent.mkdir(exist_ok=True)
            sheet.save(output, quality=95)
    write(OUT / "source-cases.private.json", case_sources)
    write(OUT / "file-inventory.private.json", file_inventory)
    write(OUT / "report-text.private.json", documents)
    write(OUT / "workbook-review.private.json", dict(path=str(path), sha256=sha(path), sheets=supplementary, media_files=embedded))
    write(OUT / "photo-map.private.json", image_manifest)
    # Render histology-named PDFs for original-page review; scans need visual reading.
    rendered = []
    for path_string, info in documents.items():
        if info.get("histology_filename"):
            pdf = pdfium.PdfDocument(path_string)
            for i in range(len(pdf)):
                output = OUT / "report-pages" / f"histology-{info['sha256'][:16]}-{i+1:02d}.png"
                output.parent.mkdir(exist_ok=True)
                pdf[i].render(scale=2.3).to_pil().save(output)
                rendered.append(dict(path=str(output), original=path_string, page=i+1, sha256=info["sha256"]))
            pdf.close()
    write(OUT / "rendered-reports.private.json", rendered)
    summary = dict(moderate_records=len(moderate), bound_folders=len({p for c in case_sources for p in c["source_folders"]}),
        source_file_types=dict(Counter(f["suffix"] for f in file_inventory)), pdf_files=len([d for d in documents.values() if "pages" in d]),
        textless_pdfs=sum(not d["text_characters"] for d in documents.values() if "pages" in d),
        histology_named_pdf_count=sum(bool(d.get("histology_filename")) for d in documents.values()),
        rendered_histology_pages=len(rendered), workbook_sheets=[dict(sheet=s["sheet"], rows=s["rows"], columns=s["columns"],
            matched_rows=len(s["matches"]), embedded_images=s["embedded_images"]) for s in supplementary],
        photo_count=len(image_manifest), photo_sheets=13, unnumbered_original_matched=sum(bool(x["selected_unnumbered_source_matches"]) for x in image_manifest),
        photo_exif_orientation_counts=dict(Counter(str(x["exif_orientation"]) for x in image_manifest)),
        extraction_errors=errors, no_original_files_modified=True, no_models_run=True, locked_test_evaluated=False,
        source_csv_sha=sha(INPUT / "sysu_samples.csv"), split_sha=sha(INPUT / "development_nested_cv_v3.json"), script_sha=sha(Path(__file__)))
    write(OUT / "preparation-summary.json", summary)
    print(json.dumps({k:v for k,v in summary.items() if k != "extraction_errors"}, ensure_ascii=True))


if __name__ == "__main__":
    main()
